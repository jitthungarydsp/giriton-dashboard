#!/usr/bin/env python3
"""Export Hub API OVERLAP rejections to a standalone Excel workbook."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = [
    "Futás ideje", "Dátum", "Raktár", "courierId", "Futár", "E-mail",
    "MűszakPro kezdés", "HUB slot", "Shift template ID", "BlockKey",
    "Egyezés", "Eltérés perc", "Reason kód", "Reason", "Serial",
]


def load_overlap_records(path: Path) -> list[dict]:
    if not path.exists():
        return []
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, list):
        return []
    return [
        row for row in payload
        if "overlap" in str(row.get("reason") or "").casefold()
    ]


def row_values(row: dict) -> list:
    reason = str(row.get("reason") or "")
    return [
        row.get("run_at") or "", row.get("work_date") or "",
        row.get("warehouse") or "", row.get("courier_id") or "",
        row.get("courier_name") or "", row.get("email") or "",
        row.get("muszakpro_shift_start") or "", row.get("hub_slot_from") or "",
        row.get("shift_template_id") or "", row.get("block_key") or "",
        row.get("match_kind") or "", row.get("match_diff_minutes") or "",
        "OVERLAP", reason, row.get("serial") or "",
    ]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    records = load_overlap_records(Path(args.input))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "OVERLAP"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C00000")
    for record in sorted(records, key=lambda x: (str(x.get("work_date")), str(x.get("warehouse")), str(x.get("muszakpro_shift_start")))):
        sheet.append(row_values(record))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 10), 36)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width

    summary = workbook.create_sheet("ÖSSZESÍTÉS")
    summary.append(["Reason kód", "Darab"])
    summary["A1"].font = summary["B1"].font = Font(bold=True)
    summary.append(["OVERLAP", len(records)])
    summary.append(["Megjegyzés", "A Hub API HTTP 200 választ adott, de a foglalást OVERLAP reason kóddal elutasította."])
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 100

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"HUB_OVERLAP_XLSX rows={len(records)} output={output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
